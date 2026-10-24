"""Fill the retained fault-line template and create a separate clickable review chain."""
from pathlib import Path
from copy import deepcopy
import html,json,re,shutil
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from .core import *
from . import figures

ASSETS=Path(__file__).parent/'assets'

def render(work,plan_path,output):
 work=Path(work).resolve();materials,sources=load(work);plan=read_json(plan_path);validate_plan(plan,materials,sources)
 output=fresh(output,[work,plan_path]);shutil.copytree(work,output/'inputs');shutil.copy2(plan_path,output/'report-plan.json')
 measures=calculate(output/'inputs',plan,materials,sources);write_json(output/'measurements.json',measures)
 cat=catalog(materials,measures);outassets=output/'figures';outassets.mkdir();proof=[];prose=[];omitted=[]
 document=Document(ASSETS/'template.docx');nodes=list(document._element.body)

 def evidence(token,text):
  eid,_,pointer_path=token.partition('#');value=pointer(cat[eid],pointer_path)
  item={'anchor':f'ref-{len(proof)+1}','text':text,'reference':token,'value':value}
  if eid.startswith('f:'):
   f=next(x for x in materials['fields'] if 'f:'+x['id']==eid);item.update(kind='原生字段',field=f['field'],source=f['source'],location=f['evidence'],unit=f['unit'],status=f['status'])
  elif eid.startswith('r:'):item.update(kind='录波原生记录',record=cat[eid]['display_name'],cfg_source=cat[eid]['cfg_source'],dat_source=cat[eid].get('dat_source'))
  else:item.update(cfg_source=cat['r:'+cat[eid]['record']]['cfg_source'],dat_source=cat['r:'+cat[eid]['record']]['dat_source'],kind='显式窗口基波计算',record=cat[eid]['record_name'],samples=cat[eid]['samples'],native_samples=cat[eid]['native_samples'],time_record_s=cat[eid]['time_record_s'],reference_channel=cat[eid]['reference'],frequency_hz=cat[eid]['frequency_hz'],quantity_conditions=[{'channel':v['name'],'unit':v['unit'],'side':v['side'],'reasons':v['reasons']} for v in cat[eid]['results']])
  proof.append(item);return item['anchor']

 def put(paragraph,text,token=None):
  if token:
   anchor=evidence(token,text);link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to('review.html#'+anchor,RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'宋体');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');rp.append(fonts);size=OxmlElement('w:sz');size.set(qn('w:val'),'21' if paragraph._p.getparent().tag==qn('w:tc') else '24');rp.append(size);run.append(rp);t=OxmlElement('w:t');t.text=text;run.append(t);link.append(run);paragraph._p.append(link)
  else:paragraph.add_run(text)

 def p(index,text,narrative=True):
  paragraph=Paragraph(nodes[index],document._body);paragraph.clear()
  for part,token in text_parts(text,cat) if narrative else [(text,None)]:put(paragraph,part,token)
  prose.append(''.join(t for t,_ in text_parts(text,cat)) if narrative else text)
  return paragraph

 def remove(index):
  node=nodes[index]
  if node.getparent() is not None:node.getparent().remove(node)

 def table(index,rows):
  t=Table(nodes[index],document._body);header=deepcopy(t.rows[0]._tr);sample=deepcopy(t.rows[1]._tr)
  for row in list(t._tbl.findall(qn('w:tr'))):t._tbl.remove(row)
  t._tbl.append(header)
  for values in rows:
   row=deepcopy(sample);t._tbl.append(row)
   for cell,(text,token) in zip(t.rows[-1].cells,values):
    par=cell.paragraphs[0];par.clear();par.paragraph_format.keep_with_next=False;put(par,str(text),token)
    for run in par.runs:
     run.font.size=Pt(10.5);run.font.name='Times New Roman';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');run.bold=False
  h=t.rows[0]._tr.get_or_add_trPr()
  if h.find(qn('w:tblHeader')) is None:h.append(OxmlElement('w:tblHeader'))
  for cell in t.rows[0].cells:
   for par in cell.paragraphs:par.paragraph_format.keep_with_next=True
  # Re-establish the third line after row count changes.
  for cell in t.rows[-1].cells:
   tcpr=cell._tc.get_or_add_tcPr();borders=tcpr.find(qn('w:tcBorders'))
   if borders is None:borders=OxmlElement('w:tcBorders');tcpr.append(borders)
   bottom=borders.find(qn('w:bottom'))
   if bottom is None:bottom=OxmlElement('w:bottom');borders.append(bottom)
   bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'12');bottom.set(qn('w:color'),'000000')
  return t

 def raw_events(end):
  r=cat['r:'+end['record']];fs=companion_fields(materials,r)
  rows=[]
  for f in fs:
   if f['field']!='protection_event' or end['protection_line'] not in (f['value'].get('name') or ''):continue
   v=f['value'];base='f:'+f['id'];time=str(v['time'] if v.get('time') is not None else v.get('raw_time') or '未声明')
   unit=v.get('unit');time+=' '+unit if unit else '（单位未声明）'
   rows.append([(str(len(rows)+1),None),(time,base+'#'),(v.get('name') or '名称未声明',base+'#'),(v.get('phase') or '未声明',base+'#'),(str(v.get('state')) if v.get('state') is not None else '未声明',base+'#')])
  return rows or [[('未取得可核对的本线路原生保护报文',None)]+[('—',None)]*4]

 def digital_rows(end):
  r=cat['r:'+end['record']];rows=[];selected={c['number']:c for c in r['digital_channels']}
  for number in end['digital_channels']:
   ch=selected[number];evs=[(i,e) for i,e in enumerate(r['events']) if e['channel']==number]
   if not evs:
    rows.append([(ch['name'],None),('记录内未变位（起始状态 '+str(ch.get('initial_state','未声明'))+'）',None),('—',None)]);continue
   cells=[]
   for i,e in evs:
    a,b=e['time_record_interval_s'];label=f"{e['from']}→{e['to']}；({a:.6g}, {b:.6g}] s"
    if not e['edge_observed']:label+='；缺口两端状态差，非观测动作沿'
    cells.append((label,'r:'+r['id']+f'#/events/{i}'))
   for i in range(0,len(cells),2):rows.append([(ch['name']+('（续）' if i else ''),None),cells[i],cells[i+1] if i+1<len(cells) else ('未观测到后续变位',None)])
  return rows or [[('未取得可核对的数字量变位',None),('—',None),('—',None)]]

 def parameters(end):
  r=cat['r:'+end['record']];fs=companion_fields(materials,r);phase=next((f['value'] for f in fs if f['field']=='fault_phase'),None)
  roles=[('故障相电压','voltage'),('故障相电流','current'),('最大零序电流','maximum_reported_zero_sequence_current'),('最大差动电流','maximum_reported_differential_current'),('故障测距','fault_distance'),('故障相别','fault_phase')];rows=[]
  for i,(label,role) in enumerate(roles,1):
   candidates=[]
   for f in fs:
    v=f['value']
    if role in {'fault_distance','fault_phase'} and f['field']==role:candidates.append(f)
    elif f['field']=='native_measurement':
     if role in {'voltage','current'} and v['quantity']==role and v['stage']=='fault' and v['phase']==phase:candidates.append(f)
     elif role=='maximum_reported_'+v['quantity'] and v['statistic']=='maximum_reported':candidates.append(f)
   if len(candidates)==1:
    f=candidates[0];v=f['value'];value=v['raw_value'] if f['field']=='native_measurement' else str(v);unit=('—' if role=='fault_phase' else f['unit'] or '未声明');token='f:'+f['id']+'#'
   else:value='当前资料未取得可核对的唯一原生值';unit='—';token=None
   rows.append([(str(i),None),(label,None),(value,token),(unit,None)])
  return rows

 # Section 1 summarizes available evidence when topology is absent.
 for n in [2,4,5,6]:remove(n)
 p(3,plan['information'][0]);last=nodes[3]
 for text in plan['information'][1:]:
  element=deepcopy(nodes[3]);last.addnext(element);par=Paragraph(element,document._body);par.clear()
  for value,token in text_parts(text,cat):put(par,value,token)
  prose.append(''.join(v for v,_ in text_parts(text,cat)));last=element

 active=[]
 for ei,end in enumerate(plan['ends']):
  if end['record'] is None:
   for start,stop in ([(14,19),(28,36),(51,64)] if ei else [(9,14),(20,28),(38,51)]):
    p(start,end['name'],False);p(start+1,'未提供该端可用原生录波，本报告不推断该端动作及波形。',False)
    for n in range(start+2,stop):remove(n)
   remove(69 if ei else 67);remove(68 if ei else 66);continue
  r=cat['r:'+end['record']];active.append(ei)
  for idx,suffix in (([(9,'保护动作报文'),(20,'开关量变位'),(38,'录波图分析')] if ei==0 else [(14,'保护动作报文'),(28,'开关量变位'),(51,'录波图分析')])):p(idx,end['name']+suffix,False)
  source=next(s for s in materials['sources'] if s['id']==r['cfg_source'])
  p(10 if ei==0 else 15,'本节采用'+r['display_name']+'对应的原生附属记录。表内保留各事件原有时间、单位、名称与状态；相对时间的起点见原文及下述说明，不与其他装置时钟直接拼接。',False)
  p(11 if ei==0 else 16,f'表{ei+1} '+plan['line']+' '+end['name']+'保护动作简报',False);table(12 if ei==0 else 17,raw_events(end));p(13 if ei==0 else 18,end['text']['messages'])
  p(21 if ei==0 else 29,'以下按本记录起点列出数字量变位区间。辅助接点与保护出口分别保留；变化区间不是主触头的精确动作时刻。',False)
  p(26 if ei==0 else 34,f'表{ei+3} '+plan['line']+' '+end['name']+'开关量变位时间顺序',False);table(27 if ei==0 else 35,digital_rows(end))
  p(39 if ei==0 else 52,'本节按所选通道的原始样本绘制波形，并以明确样本窗口计算基波量；时轴以各自记录起点为参考。',False)
  for idx,key,prefix in ([(40,'voltage','1）故障态电压特征。'),(41,'current','2）故障态电流特征。'),(42,'clearing','3）故障切除特征。'),(47,'reclose','4）重合过程特征。'),(48,'phasor','5）故障态相量特征。')] if ei==0 else [(53,'voltage','1）故障态电压特征。'),(54,'current','2）故障态电流特征。'),(55,'clearing','3）故障切除特征。'),(56,'reclose','4）重合过程特征。'),(61,'phasor','5）故障态相量特征。')]):p(idx,prefix+end['text'][key])
  p(66 if ei==0 else 68,f'表{ei+5} '+plan['line']+' '+end['name']+'保护故障信息及测距',False);table(67 if ei==0 else 69,parameters(end))
  figures.selected_csv(output/'inputs',end['record'],end['channels'],output/f'selected-samples-e{ei}.csv',materials,sources)
  for request in end['measurements']:
   if request.get('record',end['record'])!=end['record']:figures.selected_csv(output/'inputs',request['record'],end['channels'],output/f'selected-samples-e{ei}-{request["id"]}.csv',materials,sources)
 p(36,' '.join(end['name']+'：'+end['text']['digital'] for end in plan['ends'] if end['record'] is not None))
 p(65,plan['summary']['waveforms']);p(70,plan['summary']['protection']);p(71,plan['summary']['conclusion'])

 slots=[(22,0,'digital','fault'),(24,0,'digital','reclose'),(30,1,'digital','fault'),(32,1,'digital','reclose'),(43,0,'waveform','fault'),(45,0,'waveform','reclose'),(49,0,'phasor','fault'),(57,1,'waveform','fault'),(59,1,'waveform','reclose'),(62,1,'phasor','fault')]
 figure_index=[]
 for idx,ei,kind,stage in slots:
  end=plan['ends'][ei]
  if end['record'] is None:remove(idx);remove(idx+1);continue
  span=end['digital_ranges' if kind=='digital' else 'figures'].get(stage) if kind!='phasor' else None
  path=outassets/f'e{ei}-{kind}-{stage}.png';label=plan['line']+' '+end['name']+(' 开关量变位' if kind=='digital' else ' 基波相量' if kind=='phasor' else ' 电压电流波形')+(' 故障阶段' if stage=='fault' else ' 重合过程')
  exists=False
  if kind=='phasor':exists=figures.phasor(measures[f'm:e{ei}-fault'],end['channels'],path,label)
  elif span is not None:
   if kind=='digital':figures.digital(output/'inputs',end['record'],span,path,label,materials,sources,end['digital_channels'])
   else:figures.waveform(output/'inputs',end['record'],end['channels'],span,path,label,materials,sources)
   exists=True
  if not exists:remove(idx);remove(idx+1);omitted.append(label+'：该记录范围或相角条件不足，未伪造图像。');continue
  par=Paragraph(nodes[idx],document._body);par.clear();par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.keep_with_next=True;par.alignment=1
  from PIL import Image
  with Image.open(path) as image:width,height=image.size
  actual_width=min(6.25,7.4*width/height)
  par.add_run().add_picture(str(path),width=Inches(actual_width));number=len(figure_index)+1;p(idx+1,f'图{number} '+label,False).paragraph_format.keep_with_next=False
  figure_index.append({'label':label,'file':path.relative_to(output).as_posix(),'record':end['record'],'range':span,'measurement':f'm:e{ei}-fault' if kind=='phasor' else None})
 document.save(output/'故障线路报告.docx')
 write_json(output/'review.json',{'version':'0.1.0','references':proof,'figures':figure_index,'template_sha256':sha(ASSETS/'template.docx'),'omitted_figures':omitted})
 (output/'report-text.txt').write_text('\n\n'.join(prose),encoding='utf-8')
 # Source and calculation review stays separate from the report sections.
 ss={s['id']:s for s in materials['sources']}
 pages=['<!doctype html><html lang="zh"><meta charset="utf-8"><title>故障报告实际复核</title><style>body{font:16px/1.7 sans-serif;max-width:1100px;margin:30px auto;padding:20px}section{border-top:1px solid #aaa;margin-top:25px}pre{white-space:pre-wrap;overflow-wrap:anywhere}img{max-width:100%}a{color:#174da0}</style><h1>故障报告实际复核</h1><p><a href="故障线路报告.docx">打开报告</a>；报告中的参数可跳转至本页对应来源。原始文件为本机快照，无网络请求。</p>']
 for f in figure_index:pages+=['<h2>'+html.escape(f['label'])+'</h2><img src="'+html.escape(f['file'],quote=True)+'">']
 for csvfile in sorted(output.glob('selected-samples-*.csv')):
  code=csvfile.stem.removeprefix('selected-samples-');ei=int(code[1]);key='m:'+code
  name=measures[key]['record_name'] if key in measures else cat['r:'+plan['ends'][ei]['record']]['display_name']
  pages.append('<p><a href="'+csvfile.name+'">'+html.escape(name)+'：所选通道全记录样本</a></p>')
 for item in proof:
  pages.append('<section id="'+item['anchor']+'"><h2>'+html.escape(item['text'])+'</h2>')
  for key in ['source','cfg_source','dat_source']:
   if item.get(key) in ss:
    s=ss[item[key]];pages.append('<p><a href="inputs/'+html.escape(s['raw_path'],quote=True)+'">'+html.escape(s['path'])+'</a></p>')
  pages.append('<pre>'+html.escape(json.dumps(item,ensure_ascii=False,indent=2))+'</pre></section>')
 pages.append('</html>');(output/'review.html').write_text('\n'.join(pages),encoding='utf-8')
 (output/'待补资料.md').write_text('# 具体资料与适用限制\n\n'+('\n'.join('- '+s for s in omitted) or '无因图位条件导致的删图。')+'\n\n'+ '\n'.join('- '+r['display_name']+'：'+r['error'] for r in materials['records'] if 'error' in r)+'\n\n原生数值未取得、时间基准、测点、定值和方向判断的限制已在报告相应正文说明。\n',encoding='utf-8')
 write_json(output/'checksums.json',{p.relative_to(output).as_posix():sha(p) for p in sorted(output.rglob('*')) if p.is_file()})
 checked=verify_report(output)
 if not checked['ok']:raise CaseError('Generated report failed verification: '+str(checked))
 return {'report':str(output/'故障线路报告.docx'),'review':str(output/'review.html'),'figures':len(figure_index),'references':len(proof),'verification':checked}

def verify_report(output):
 try:
  output=Path(output).resolve();checks=read_json(output/'checksums.json');actual={p.relative_to(output).as_posix() for p in output.rglob('*') if p.is_file() and p.name!='checksums.json'}
  if actual!=set(checks):raise CaseError('Report file set differs')
  for rel,h in checks.items():
   p=(output/rel).resolve()
   if not p.is_relative_to(output) or sha(p)!=h:raise CaseError('Report file changed: '+rel)
  materials,sources=load(output/'inputs');plan=read_json(output/'report-plan.json');validate_plan(plan,materials,sources)
  if calculate(output/'inputs',plan,materials,sources)!=read_json(output/'measurements.json'):raise CaseError('Phasor replay differs')
  if read_json(output/'review.json')['template_sha256']!=sha(ASSETS/'template.docx'):raise CaseError('Template differs; use the package matching this report')
  cat=catalog(materials,read_json(output/'measurements.json'))
  for r in read_json(output/'review.json')['references']:
   eid,_,path=r['reference'].partition('#')
   if pointer(cat[eid],path)!=r['value']:raise CaseError('Source reference changed: '+r['reference'])
  with ZipFile(output/'故障线路报告.docx') as z:
   from xml.etree import ElementTree as E
   relations=E.fromstring(z.read('word/_rels/document.xml.rels'));targets={r.get('Target') for r in relations if r.get('Type','').endswith('/hyperlink')}
   expected={'review.html#'+r['anchor'] for r in read_json(output/'review.json')['references']}
   if targets!=expected:raise CaseError('Word parameter links differ from review anchors')
   review_html=(output/'review.html').read_text(encoding='utf-8')
   if any('id="'+r['anchor']+'"' not in review_html for r in read_json(output/'review.json')['references']):raise CaseError('Review anchor missing')
   root=E.fromstring(z.read('word/document.xml'));text=''.join(x.text or '' for x in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
   if re.search(r'〈|〉|\[\[TODO|【待补',text):raise CaseError('Placeholder in finished report')
   for title in ['故障信息','录波图解析','综合总结']:
    if title not in text:raise CaseError('Template section missing: '+title)
  return {'ok':True,'files_checked':len(checks),'meaning':'Source/window/document integrity and reproducible quantities; electrical reasoning and layout are reviewed separately'}
 except (OSError,CaseError,ValueError,KeyError,TypeError,IndexError) as exc:return {'ok':False,'error':str(exc)}
